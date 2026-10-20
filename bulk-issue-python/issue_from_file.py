"""Emisión masiva de DTEs desde un CSV o Excel con el SDK de Tupana.

Cada fila es una línea de detalle; las filas con el mismo ``ref`` forman un
documento. Se envía en lotes de hasta 200 documentos con un Idempotency-Key
derivado del contenido: si el script se corta y lo vuelves a correr con el
mismo archivo, la API devuelve el lote original en vez de emitir de nuevo.

Por defecto NO emite nada (modo prueba): valida el archivo y muestra el resumen.
Agrega ``--emit`` para emitir de verdad.

    export TUPANA_API_KEY=...
    python issue_from_file.py invoices.csv --issuer-rut 76543210-3            # prueba
    python issue_from_file.py invoices.xlsx --issuer-rut 76543210-3 --emit    # emite
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import sys
from collections import OrderedDict
from collections.abc import Iterable, Iterator
from pathlib import Path
from typing import Any

from tupana import MAX_BATCH_SIZE, APIError, Tupana

REQUIRED_COLUMNS = {"ref", "dte_type", "receiver_rut", "item_name", "quantity", "unit_price"}


def read_rows(path: Path) -> list[dict[str, str]]:
    """Lee un .csv (UTF-8, coma o punto y coma) o un .xlsx (primera hoja)."""
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows)]
        return [
            {h: "" if v is None else str(v).strip() for h, v in zip(header, row, strict=False)}
            for row in rows
            if any(v not in (None, "") for v in row)
        ]
    with path.open(newline="", encoding="utf-8-sig") as fh:
        sample = fh.read(2048)
        fh.seek(0)
        dialect = csv.Sniffer().sniff(sample, delimiters=",;")
        return [
            {k.strip(): (v or "").strip() for k, v in row.items() if k}
            for row in csv.DictReader(fh, dialect=dialect)
        ]


_CHILEAN_NUMBER = re.compile(r"^-?\d{1,3}(\.\d{3})+(,\d+)?$")


def _number(value: str, field: str, line: int) -> float:
    """Acepta 49990, 49990.5, 49.990 y 49.990,5 (formato chileno)."""
    text = value.replace("$", "").replace(" ", "")
    if _CHILEAN_NUMBER.match(text) or ("," in text and "." not in text):
        text = text.replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        raise SystemExit(f"Fila {line}: '{field}' no es un número ({value!r})") from None


def clp(amount: float) -> str:
    return "$" + f"{amount:,.0f}".replace(",", ".")


def build_documents(rows: Iterable[dict[str, str]], issuer_rut: str) -> list[dict[str, Any]]:
    missing = REQUIRED_COLUMNS - set(next(iter(rows), {}).keys())
    if missing:
        raise SystemExit(f"Faltan columnas: {', '.join(sorted(missing))}")
    docs: OrderedDict[str, dict[str, Any]] = OrderedDict()
    for line, row in enumerate(rows, start=2):
        ref = row["ref"]
        doc = docs.get(ref)
        if doc is None:
            doc = {
                "dte_type": {"code": row["dte_type"]},
                "document_issuer": {"rut": issuer_rut},
                "document_receiver": {"rut": row["receiver_rut"]},
                "details": [],
                # Tu propio id: permite reencontrar el documento con
                # documents.list(..., json_param={"external_ref": ref}).
                "json_param": {"external_ref": ref},
            }
            if row.get("due_date"):
                doc["header"] = {"due_date": row["due_date"]}
            docs[ref] = doc
        elif doc["document_receiver"]["rut"] != row["receiver_rut"]:
            raise SystemExit(f"Fila {line}: el documento {ref} tiene dos receptores distintos")
        item: dict[str, Any] = {
            "item_name": row["item_name"],
            "quantity": _number(row["quantity"], "quantity", line),
            "unit_price": _number(row["unit_price"], "unit_price", line),
        }
        if row.get("item_description"):
            item["item_description"] = row["item_description"]
        doc["details"].append(item)
    return list(docs.values())


def chunks(
    items: list[dict[str, Any]], size: int = MAX_BATCH_SIZE
) -> Iterator[list[dict[str, Any]]]:
    for start in range(0, len(items), size):
        yield items[start : start + size]


def idempotency_key(chunk: list[dict[str, Any]]) -> str:
    """Misma entrada → misma key → la API no vuelve a emitir el lote."""
    digest = hashlib.sha256(json.dumps(chunk, sort_keys=True).encode()).hexdigest()
    return f"bulk-{digest[:48]}"


def summarize(docs: list[dict[str, Any]]) -> None:
    for doc in docs:
        net = sum(d["quantity"] * d["unit_price"] for d in doc["details"])
        print(
            f"  {doc['json_param']['external_ref']:<12} DTE {doc['dte_type']['code']:<3} "
            f"→ {doc['document_receiver']['rut']:<12} {len(doc['details'])} línea(s), "
            f"neto {clp(net)}"
        )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawTextHelpFormatter
    )
    parser.add_argument("file", type=Path, help="CSV o XLSX")
    parser.add_argument("--issuer-rut", default=os.environ.get("TUPANA_ISSUER_RUT"))
    parser.add_argument("--emit", action="store_true", help="emite de verdad (sin esto es prueba)")
    parser.add_argument("--results", type=Path, default=Path("results.csv"))
    args = parser.parse_args(argv)
    if not args.issuer_rut:
        parser.error("indica --issuer-rut o TUPANA_ISSUER_RUT")

    docs = build_documents(read_rows(args.file), args.issuer_rut)
    print(f"{len(docs)} documento(s) en {args.file}:")
    summarize(docs)
    if not args.emit:
        print("\nModo prueba: no se emitió nada. Agrega --emit para emitir.")
        return 0

    tupana = Tupana()
    results: list[dict[str, Any]] = []
    for number, chunk in enumerate(chunks(docs), start=1):
        print(f"\nLote {number}: enviando {len(chunk)} documento(s)…")
        try:
            created = tupana.documents.create_batch(
                chunk, use_defaults=True, idempotency_key=idempotency_key(chunk)
            )
        except APIError as err:
            print(f"  Error {err.status_code} ({err.code}): {err.message}", file=sys.stderr)
            return 1
        for warning in created.get("warnings") or []:
            print(f"  Aviso doc {warning.get('document_index')}: {warning.get('message')}")
        batch = tupana.documents.wait_for_batch(created["batch_id"], timeout=600)
        for item in batch.get("documents") or []:
            doc = item.get("document") or {}
            ref = chunk[item["index"]]["json_param"]["external_ref"]
            results.append(
                {
                    "ref": ref,
                    "status": item.get("status"),
                    "folio": doc.get("folio"),
                    "document_id": doc.get("id"),
                    "pdf_url": doc.get("pdf_url"),
                }
            )
            print(f"  {ref:<12} {item.get('status'):<16} folio {doc.get('folio')}")

    with args.results.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=["ref", "status", "folio", "document_id", "pdf_url"])
        writer.writeheader()
        writer.writerows(results)
    print(f"\nResultados en {args.results}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
